"""
CareerCompass AI — Official Java PBL Report Generator
Generates submission-ready Java PBL Report DOCX adhering strictly to:
- Official CIT Java PBL Report Template structure
- 100% Real Java Codebase Evidence (Spring Boot 3.2.0, Java 17/21, JPA, JJWT)
- Exact institutional front matter, certificate, declaration, and acknowledgements
- Real code snippets, architecture diagram, screenshots, and test results
"""

import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_TEMPLATE_PATH = r'C:\Users\Sabarish\Downloads\CareerCompass_AI_PBL_Report.docx'
TARGET_DOCX_PATH = r'C:\Users\Sabarish\Downloads\CareerCompass_AI_Java_PBL_Report.docx'
PROJECT_DOCX_PATH = r'c:\sabarish\Projects\Career Compass\docs\CareerCompass_AI_Java_PBL_Report.docx'
IMAGE_DIR = r'c:\sabarish\Projects\Career Compass\scratch\java_report_images'

def set_cell_properties(cell, fill_hex=None, bold=False, font_size_pt=9.0, align=WD_ALIGN_PARAGRAPH.LEFT, text=""):
    cell.text = text
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    if p.runs:
        r = p.runs[0]
        r.font.name = "Calibri"
        r.font.size = Pt(font_size_pt)
        r.bold = bold
        if fill_hex:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
            tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.25)
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>')
    tcPr.append(shd)
    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
    
    cell.text = code_text
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    for r in p.runs:
        r.font.name = "Consolas"
        r.font.size = Pt(8.0)
        r.font.color.rgb = RGBColor(0x21, 0x21, 0x21)

print("Java Report Generator module loaded.")
